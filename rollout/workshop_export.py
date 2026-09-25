"""The outcome of a Fit-to-Standard workshop, as a file the client keeps.

One outcome, four formats: Markdown, PDF, Word and Excel. Every format is
drawn from the same `outcome()` rows, so they cannot disagree about what was
decided -- the PDF goes further and is rendered from the Markdown, as the
workshop pack is.

Two scopes. With a session id it is that sitting of facilitator mode: what the
room answered when it pressed Submit. Without one it is the run's standing
record: the current decision on every gap, however and whenever it was made.
"""

from __future__ import annotations

import datetime as _dt
import io
import re

from .decisions import LETTERS

FORMATS = {
    "md": ("text/markdown; charset=utf-8", "md"),
    "pdf": ("application/pdf", "pdf"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
}
VERDICT = {"accept": "Accepted", "defer": "Deferred", "reject": "Rejected"}
COLUMNS = ["#", "Gap", "Question", "Decision", "Option chosen", "Rationale",
           "Type", "Materiality", "As-Is step", "Decision owners", "Decided by", "Decided at"]


def _when(iso: str | None) -> str:
    if not iso:
        return ""
    try:
        return _dt.datetime.fromisoformat(iso).strftime("%d %b %Y %H:%M")
    except ValueError:
        return iso


def outcome(run: dict, decisions: list[dict], sessions: list[dict], session_id: str = "") -> dict:
    """The rows and the facts around them, in the agenda's running order.

    Only current decisions count for the run-wide record: a verdict that was
    later replaced is history, not the outcome. For one session, its own rows
    are the outcome, whatever came after."""
    session = next((s for s in sessions if s["id"] == session_id), None) if session_id else None
    if session_id and not session:
        raise LookupError(session_id)
    chosen = ([d for d in decisions if d.get("session_id") == session_id] if session
              else [d for d in decisions if d.get("is_current", True)])
    by_gap: dict[str, dict] = {}
    for d in sorted(chosen, key=lambda x: (x.get("decided_at") or "", x.get("id") or 0)):
        by_gap[d["gap_id"]] = d

    agenda = (run.get("scores") or {}).get("agenda") or []
    order = [a["gap_id"] for a in agenda] + sorted(g for g in by_gap if g not in {a["gap_id"] for a in agenda})
    rows, still_open = [], []
    for gap in order:
        d = by_gap.get(gap)
        if not d:
            item = next(a for a in agenda if a["gap_id"] == gap)
            still_open.append({"gap_id": gap, "question": item.get("topic", "")})
            continue
        option = ""
        if d.get("option_index") is not None and d.get("option_text"):
            option = f"{LETTERS[d['option_index']]}: {d['option_text']}"
        rows.append({
            "#": len(rows) + 1, "Gap": gap, "Question": d.get("question", ""),
            "Decision": VERDICT.get(d["verdict"], d["verdict"]), "Option chosen": option,
            "Rationale": d.get("rationale", ""), "Type": d.get("primary_type", ""),
            "Materiality": d.get("materiality", ""), "As-Is step": d.get("as_is_step_id", ""),
            "Decision owners": ", ".join(d.get("decision_owner") or []),
            "Decided by": d.get("reviewer", ""), "Decided at": _when(d.get("decided_at")),
        })

    counts = {v: sum(1 for r in rows if r["Decision"] == v) for v in VERDICT.values()}
    analysis = run.get("analysis") or {}
    title = run.get("scope_label") or (analysis.get("template_process") or "").split(" (")[0] or "Fit-to-Standard"
    facts = [("Process", title), ("Country", run.get("country") or "—"),
             ("Global Template process", analysis.get("template_process") or "—"),
             ("Analysis", run.get("id", ""))]
    if session:
        facts += [("Facilitator", session.get("facilitator") or "—"),
                  ("In the room", ", ".join(session.get("attendees") or []) or "—"),
                  ("Submitted", _when(session.get("submitted_at") or session.get("started_at")))]
    else:
        facts += [("Workshop sittings", str(len(sessions)) if sessions else "—")]
    facts += [("Decisions", f"{len(rows)} — {counts['Accepted']} accepted, "
                            f"{counts['Deferred']} deferred, {counts['Rejected']} rejected"),
              ("Still open", str(len(still_open)))]
    return {"title": f"Workshop outcome — {title}" + (f" · {run['country']}" if run.get("country") else ""),
            "scope": "This sitting" if session else "Current decisions on every gap",
            "facts": facts, "rows": rows, "open": still_open,
            "generated": _dt.datetime.now().strftime("%d %B %Y at %H:%M")}


def filename(run: dict, fmt: str, session_id: str = "") -> str:
    bits = [b for b in (run.get("country"), run.get("scope_label")) if b]
    stem = " - ".join(str(b) for b in bits) or str(run.get("id") or "workshop")
    stem = re.sub(r'[\\/:*?"<>|]+', " ", stem)
    stem = re.sub(r"\s+", " ", stem.encode("ascii", "ignore").decode()).strip()[:100]
    return f"Workshop outcome - {stem}{' - ' + session_id if session_id else ''}.{FORMATS[fmt][1]}"


def _cell(text) -> str:
    return str(text or "").replace("|", "\\|").replace("\n", " ").strip()


def to_markdown(o: dict) -> str:
    out = [f"# {o['title']}", "", f"*{o['scope']}. Generated {o['generated']}.*", ""]
    out += ["| | |", "|---|---|"] + [f"| **{k}** | {_cell(v)} |" for k, v in o["facts"]] + [""]
    out += ["## Decisions", ""]
    if o["rows"]:
        cols = [c for c in COLUMNS if c != "#"]
        out += ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
        out += ["| " + " | ".join(_cell(r[c]) for c in cols) + " |" for r in o["rows"]]
    else:
        out.append("No decisions recorded.")
    out.append("")
    if o["open"]:
        out += ["## Still open", ""] + [f"- **{x['gap_id']}** — {x['question']}" for x in o["open"]] + [""]
    return "\n".join(out)


def to_pdf(run: dict, o: dict) -> bytes:
    from . import pdf
    return pdf.render(run, to_markdown(o))


def to_docx(o: dict) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    # Landscape: the decisions table is wide, and a portrait page squeezes
    # the question and the rationale into columns a word wide.
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.6))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading(o["title"], level=1)
    doc.add_paragraph(f"{o['scope']}. Generated {o['generated']}.").runs[0].italic = True
    facts = doc.add_table(rows=0, cols=2)
    facts.style = "Light List Accent 1"
    for k, v in o["facts"]:
        cells = facts.add_row().cells
        cells[0].text, cells[1].text = k, str(v)
        cells[0].paragraphs[0].runs[0].bold = True

    doc.add_heading("Decisions", level=2)
    cols = [c for c in COLUMNS if c not in ("As-Is step", "Decided at")]
    if o["rows"]:
        t = doc.add_table(rows=1, cols=len(cols))
        t.style = "Light Grid Accent 1"
        for cell, name in zip(t.rows[0].cells, cols):
            cell.text = name
        for r in o["rows"]:
            for cell, name in zip(t.add_row().cells, cols):
                cell.text = str(r[name])
        for row in t.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(8.5)
    else:
        doc.add_paragraph("No decisions recorded.")
    if o["open"]:
        doc.add_heading("Still open", level=2)
        for x in o["open"]:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(x["gap_id"]).bold = True
            p.add_run(f" — {x['question']}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_xlsx(o: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Decisions"
    ws.append(COLUMNS)
    for r in o["rows"]:
        ws.append([r[c] for c in COLUMNS])
    head = PatternFill("solid", fgColor="0E6E68")
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), head
    widths = {"#": 5, "Gap": 16, "Question": 60, "Decision": 12, "Option chosen": 45, "Rationale": 45,
              "Type": 8, "Materiality": 12, "As-Is step": 11, "Decision owners": 30,
              "Decided by": 18, "Decided at": 18}
    for i, c in enumerate(COLUMNS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths[c]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Summary")
    info.append([o["title"]]); info["A1"].font = Font(bold=True, size=13)
    info.append([o["scope"]]); info.append([f"Generated {o['generated']}"]); info.append([])
    for k, v in o["facts"]:
        info.append([k, v]); info.cell(info.max_row, 1).font = Font(bold=True)
    if o["open"]:
        info.append([]); info.append(["Still open"]); info.cell(info.max_row, 1).font = Font(bold=True)
        for x in o["open"]:
            info.append([x["gap_id"], x["question"]])
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 26, 90
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render(run: dict, o: dict, fmt: str) -> bytes:
    if fmt == "md":
        return to_markdown(o).encode()
    if fmt == "pdf":
        return to_pdf(run, o)
    if fmt == "docx":
        return to_docx(o)
    if fmt == "xlsx":
        return to_xlsx(o)
    raise ValueError(f"Unknown format {fmt}")
