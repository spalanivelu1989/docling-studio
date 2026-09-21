"""The BPML process hierarchy: the Copilot's scope backbone (handover §2).

`BPML_ProcessesHierarchyExtended.xlsx` is read directly rather than through
the corpus, because its Markdown conversion is a stub -- 9,096 rows x 50
columns is "too wide to render as a table", so the hierarchy exists in the
index as a heading and nothing else. See fitgap/NOTES.md.
"""

from __future__ import annotations

import re
import threading
from dataclasses import dataclass, field
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
SHEET = BASE_DIR / "solvay-spark" / "pkg" / "BPML_ProcessesHierarchyExtended.xlsx"

# Rows start under a four-line banner; the header row names the columns.
HEADER_ROW = 5
NAME_COL = "Process Name (1033)"
DESC_COL = "Description (1033)"

# "4.5.1.4  Create Sales Order" -> code, name. A level-1 process is written
# "4.0"; every deeper level drops the trailing zero and adds a segment.
_LINE = re.compile(r"^(\d+(?:\.\d+)+)\s+(.*)$")

STREAM_OF_ROOT = {
    "1": "H2R", "2": "A2D", "4": "L2C", "5": "F2S",
    "6": "P2P", "7": "P2P", "8": "I2D", "9": "R2R",
}


def level_of(code: str) -> int:
    """4.0 -> 1, 4.5 -> 2, 4.5.1 -> 3, 4.5.1.4 -> 4."""
    parts = code.split(".")
    if len(parts) == 2 and parts[1] == "0":
        return 1
    return len(parts)


def parent_of(code: str) -> str | None:
    parts = code.split(".")
    if level_of(code) == 1:
        return None
    if len(parts) == 2:           # 4.5 -> 4.0
        return f"{parts[0]}.0"
    return ".".join(parts[:-1])   # 4.5.1.4 -> 4.5.1


@dataclass
class Process:
    code: str
    name: str
    level: int
    parent: str | None
    description: str = ""
    process_type: str = ""
    status: str = ""
    children: list[str] = field(default_factory=list)

    @property
    def stream(self) -> str | None:
        return STREAM_OF_ROOT.get(self.code.split(".")[0])

    def brief(self) -> dict:
        return {"code": self.code, "name": self.name, "level": self.level, "parent": self.parent}

    def full(self) -> dict:
        return {**self.brief(), "description": self.description, "process_type": self.process_type,
                "status": self.status, "children": self.children, "stream": self.stream}


_lock = threading.Lock()
_cache: dict[str, Process] | None = None
_load_error: str | None = None


def load(force: bool = False) -> dict[str, Process]:
    """Parse the sheet once per process. Later duplicate codes are dropped:
    the sheet carries a second "2.0 A2D" row that would otherwise overwrite
    "2.0 Acquire to Dispose"."""
    global _cache, _load_error
    with _lock:
        if _cache is not None and not force:
            return _cache
        procs: dict[str, Process] = {}
        try:
            import openpyxl

            wb = openpyxl.load_workbook(SHEET, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            rows = ws.iter_rows(values_only=True)
            header: list[str] = []
            for n, row in enumerate(rows, 1):
                if n == HEADER_ROW:
                    header = [str(c).strip() if c else "" for c in row]
                    break
            name_i = header.index(NAME_COL) if NAME_COL in header else 2
            desc_i = header.index(DESC_COL) if DESC_COL in header else 4
            type_i = header.index("Process Type") if "Process Type" in header else 15
            stat_i = header.index("Status") if "Status" in header else 17

            def cell(row, i):
                return str(row[i]).strip() if i < len(row) and row[i] is not None else ""

            for row in rows:
                m = _LINE.match(cell(row, name_i))
                if not m:
                    continue
                code, name = m.group(1), m.group(2).strip()
                if code in procs:
                    continue
                procs[code] = Process(
                    code=code, name=name, level=level_of(code), parent=parent_of(code),
                    description=cell(row, desc_i)[:1500],
                    process_type=cell(row, type_i), status=cell(row, stat_i),
                )
            wb.close()
            _load_error = None
        except Exception as exc:  # a missing sheet must not take the API down
            _load_error = f"{type(exc).__name__}: {exc}"

        for code, p in procs.items():
            if p.parent and p.parent in procs:
                procs[p.parent].children.append(code)
        for p in procs.values():
            p.children.sort(key=sort_key)

        _cache = procs
        return _cache


def load_error() -> str | None:
    load()
    return _load_error


def sort_key(code: str) -> tuple:
    """Numeric sort, so 4.10 follows 4.9 instead of 4.1."""
    return tuple(int(x) if x.isdigit() else 0 for x in code.split("."))


def get(code: str) -> Process | None:
    return load().get(code.strip())


def exists(code: str) -> bool:
    return code.strip() in load()


def roots() -> list[Process]:
    return [p for p in sorted(load().values(), key=lambda p: sort_key(p.code)) if p.level == 1]


def children(code: str) -> list[Process]:
    p = get(code)
    procs = load()
    return [procs[c] for c in p.children] if p else []


def subtree(code: str, max_depth: int | None = None) -> list[Process]:
    """`code` first, then every descendant, depth first, in code order."""
    root = get(code)
    if not root:
        return []
    out: list[Process] = []
    stack = [root]
    while stack:
        p = stack.pop(0)
        out.append(p)
        if max_depth is not None and p.level - root.level >= max_depth:
            continue
        stack = children(p.code) + stack
    return out


def steps_in_scope(code: str, max_steps: int | None = None) -> list[Process]:
    """The units the Copilot classifies: level-4 steps under `code`, falling
    back to level 3 (then the node itself) where level 4 does not exist.

    The fallback is per branch, not per scope: 4.5 may detail some of its
    level-3 steps down to level 4 and leave others at level 3, and dropping
    the latter would silently shrink the register.
    """
    root = get(code)
    if not root:
        return []
    picked: list[Process] = []

    def walk(p: Process) -> None:
        kids = children(p.code)
        if p.level >= 4 or not kids:
            picked.append(p)
            return
        for k in kids:
            walk(k)

    walk(root)
    picked.sort(key=lambda p: sort_key(p.code))
    return picked[:max_steps] if max_steps else picked


def search(text: str, limit: int = 12) -> list[Process]:
    """Resolve what the user typed to processes: a code, a code prefix, or
    words from a name. Ordered best match first."""
    q = (text or "").strip().lower()
    if not q:
        return []
    procs = load()
    if q in procs:
        return [procs[q]]

    scored: list[tuple[float, Process]] = []
    words = [w for w in re.split(r"[^a-z0-9.]+", q) if len(w) > 2]
    for p in procs.values():
        name = p.name.lower()
        score = 0.0
        if p.code.startswith(q):
            score += 6.0
        if q in name:
            score += 5.0
        hits = sum(1 for w in words if w in name or w in p.code)
        if hits:
            score += 2.0 * hits / max(len(words), 1) + 0.6 * hits
        if score:
            # Prefer the shallower, more quotable process when two tie.
            score -= 0.15 * p.level
            scored.append((score, p))
    scored.sort(key=lambda s: (-s[0], sort_key(s[1].code)))
    return [p for _, p in scored[:limit]]


def resolve_scope(text: str) -> Process | None:
    """One process to run a register over, from a code or a free-text phrase."""
    hits = search(text)
    return hits[0] if hits else None


def stats() -> dict:
    procs = load()
    by_level: dict[int, int] = {}
    for p in procs.values():
        by_level[p.level] = by_level.get(p.level, 0) + 1
    return {
        "sheet": str(SHEET.relative_to(BASE_DIR)),
        "available": SHEET.is_file() and not _load_error,
        "error": _load_error,
        "processes": len(procs),
        "by_level": dict(sorted(by_level.items())),
        "roots": [p.brief() for p in roots()],
    }
